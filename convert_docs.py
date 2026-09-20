import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Preformatted, KeepTogether, HRFlowable
from reportlab.lib import colors

def parse_markdown_to_structure(md_text):
    """
    Parses simple markdown into structured elements:
    - ('h1', text)
    - ('h2', text)
    - ('h3', text)
    - ('p', text)
    - ('bullet', text)
    - ('code_block', code_text)
    - ('table', header_list, row_list)
    - ('hr', '')
    """
    lines = md_text.split('\n')
    elements = []
    i = 0
    in_code_block = False
    code_lines = []
    
    while i < len(lines):
        line = lines[i]
        
        # Code block toggle
        if line.strip().startswith('```'):
            if in_code_block:
                elements.append(('code_block', '\n'.join(code_lines)))
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue
            
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue
            
        line_stripped = line.strip()
        
        if not line_stripped:
            i += 1
            continue
            
        if line_stripped == '---' or line_stripped == '***':
            elements.append(('hr', ''))
            i += 1
            continue
            
        # Table detection
        if line_stripped.startswith('|') and line_stripped.endswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
                
            if len(table_lines) >= 2:
                # Parse table
                raw_headers = [c.strip() for c in table_lines[0].strip('|').split('|')]
                raw_rows = []
                for row_line in table_lines[1:]:
                    if re.match(r'^\|[\s\-:|]+\|$', row_line):
                        continue # separator
                    raw_cells = [c.strip() for c in row_line.strip('|').split('|')]
                    raw_rows.append(raw_cells)
                elements.append(('table', raw_headers, raw_rows))
            continue
            
        # Headings
        if line_stripped.startswith('# '):
            elements.append(('h1', line_stripped[2:].strip()))
            i += 1
            continue
        elif line_stripped.startswith('## '):
            elements.append(('h2', line_stripped[3:].strip()))
            i += 1
            continue
        elif line_stripped.startswith('### '):
            elements.append(('h3', line_stripped[4:].strip()))
            i += 1
            continue
            
        # Bullet list
        if line_stripped.startswith('- ') or line_stripped.startswith('* '):
            elements.append(('bullet', line_stripped[2:].strip()))
            i += 1
            continue
            
        # Numbered list
        num_match = re.match(r'^\d+\.\s+(.*)$', line_stripped)
        if num_match:
            elements.append(('bullet', num_match.group(1).strip()))
            i += 1
            continue
            
        # Standard paragraph
        elements.append(('p', line_stripped))
        i += 1
        
    return elements

def clean_md_text(text):
    # Remove markdown link formatting [text](url) -> text (url)
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', text)
    # Remove bold/italic markers for clean text
    text = text.replace('**', '').replace('*', '').replace('`', '')
    return text

def create_word_doc(elements, title, output_path):
    doc = docx.Document()
    
    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    for el in elements:
        el_type = el[0]
        
        if el_type == 'h1':
            p = doc.add_paragraph()
            run = p.add_run(clean_md_text(el[1]))
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = RGBColor(14, 43, 92)
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            
        elif el_type == 'h2':
            p = doc.add_paragraph()
            run = p.add_run(clean_md_text(el[1]))
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(27, 73, 140)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            
        elif el_type == 'h3':
            p = doc.add_paragraph()
            run = p.add_run(clean_md_text(el[1]))
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(40, 90, 160)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            
        elif el_type == 'p':
            p = doc.add_paragraph()
            run = p.add_run(clean_md_text(el[1]))
            run.font.size = Pt(10.5)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            
        elif el_type == 'bullet':
            p = doc.add_paragraph(style='List Bullet')
            run = p.add_run(clean_md_text(el[1]))
            run.font.size = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            
        elif el_type == 'code_block':
            p = doc.add_paragraph()
            run = p.add_run(el[1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(30, 30, 30)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            
        elif el_type == 'hr':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            
        elif el_type == 'table':
            headers = el[1]
            rows = el[2]
            if not headers:
                continue
            table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            
            # Style header
            for c_idx, head_text in enumerate(headers):
                cell = table.cell(0, c_idx)
                cell.text = clean_md_text(head_text)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True
                        run.font.size = Pt(9.5)
                        run.font.color.rgb = RGBColor(255, 255, 255)
                # Background color for header
                shd = parse_xml(r'<w:shd {} w:fill="1E3A8A"/>'.format(nsdecls('w')))
                cell._tc.get_or_add_tcPr().append(shd)
                
            # Style rows
            for r_idx, row_data in enumerate(rows):
                for c_idx, cell_text in enumerate(row_data):
                    if c_idx < len(headers):
                        cell = table.cell(r_idx + 1, c_idx)
                        cell.text = clean_md_text(cell_text)
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                run.font.size = Pt(9)
                        if r_idx % 2 == 1:
                            shd = parse_xml(r'<w:shd {} w:fill="F1F5F9"/>'.format(nsdecls('w')))
                            cell._tc.get_or_add_tcPr().append(shd)
            p_after = doc.add_paragraph()
            p_after.paragraph_format.space_after = Pt(6)
            
    doc.save(output_path)
    print(f"Created Word document: {output_path}")

def create_pdf_doc(elements, title, output_path):
    pdf = SimpleDocTemplate(
        output_path,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    
    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#0E2B5C'),
        spaceAfter=10,
        fontName='Helvetica-Bold'
    )
    
    h2_style = ParagraphStyle(
        'DocH2',
        parent=styles['Heading2'],
        fontSize=13,
        leading=16,
        textColor=colors.HexColor('#1B498C'),
        spaceBefore=10,
        spaceAfter=6,
        fontName='Helvetica-Bold'
    )
    
    h3_style = ParagraphStyle(
        'DocH3',
        parent=styles['Heading3'],
        fontSize=10.5,
        leading=13,
        textColor=colors.HexColor('#285AA0'),
        spaceBefore=7,
        spaceAfter=3,
        fontName='Helvetica-Bold'
    )
    
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontSize=9,
        leading=12.5,
        textColor=colors.HexColor('#1E293B'),
        spaceAfter=4
    )
    
    bullet_style = ParagraphStyle(
        'DocBullet',
        parent=styles['Normal'],
        fontSize=8.5,
        leading=11.5,
        leftIndent=12,
        textColor=colors.HexColor('#1E293B'),
        spaceAfter=2
    )
    
    code_style = ParagraphStyle(
        'DocCode',
        parent=styles['Normal'],
        fontName='Courier',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor('#0F172A')
    )
    
    story = []
    
    for el in elements:
        el_type = el[0]
        
        if el_type == 'h1':
            story.append(Paragraph(clean_md_text(el[1]), title_style))
        elif el_type == 'h2':
            story.append(Paragraph(clean_md_text(el[1]), h2_style))
        elif el_type == 'h3':
            story.append(Paragraph(clean_md_text(el[1]), h3_style))
        elif el_type == 'p':
            story.append(Paragraph(clean_md_text(el[1]), body_style))
        elif el_type == 'bullet':
            story.append(Paragraph(f"&bull; {clean_md_text(el[1])}", bullet_style))
        elif el_type == 'hr':
            story.append(Spacer(1, 4))
            story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#CBD5E1'), spaceAfter=6, spaceBefore=4))
        elif el_type == 'code_block':
            escaped_code = el[1].replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
            story.append(Spacer(1, 3))
            story.append(Preformatted(escaped_code, code_style))
            story.append(Spacer(1, 3))
        elif el_type == 'table':
            headers = el[1]
            rows = el[2]
            if not headers:
                continue
            
            table_data = []
            hdr_cells = [Paragraph(f"<b>{clean_md_text(h)}</b>", ParagraphStyle('TH', parent=body_style, textColor=colors.white, fontSize=8, leading=10)) for h in headers]
            table_data.append(hdr_cells)
            
            for row in rows:
                row_cells = []
                for idx, cell in enumerate(row):
                    if idx < len(headers):
                        row_cells.append(Paragraph(clean_md_text(cell), ParagraphStyle('TD', parent=body_style, fontSize=7.5, leading=9.5)))
                table_data.append(row_cells)
                
            col_count = len(headers)
            available_width = 540
            col_width = available_width / col_count
            
            t = Table(table_data, colWidths=[col_width] * col_count)
            t.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
                ('TOPPADDING', (0, 0), (-1, -1), 4),
                ('LEFTPADDING', (0, 0), (-1, -1), 4),
                ('RIGHTPADDING', (0, 0), (-1, -1), 4),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ]))
            story.append(Spacer(1, 4))
            story.append(t)
            story.append(Spacer(1, 6))
            
    pdf.build(story)
    print(f"Created PDF document: {output_path}")

def main():
    workspace = os.path.dirname(os.path.abspath(__file__))
    
    # 1. Project Documentation
    doc1_md_path = os.path.join(workspace, 'PROJECT_DOCUMENTATION.md')
    if os.path.exists(doc1_md_path):
        with open(doc1_md_path, 'r', encoding='utf-8') as f:
            md1 = f.read()
        elements1 = parse_markdown_to_structure(md1)
        create_word_doc(elements1, "AI-NIDS Project Documentation", os.path.join(workspace, 'PROJECT_DOCUMENTATION.docx'))
        create_pdf_doc(elements1, "AI-NIDS Project Documentation", os.path.join(workspace, 'PROJECT_DOCUMENTATION.pdf'))
        
    # 2. Installation & Run Guide
    doc2_md_path = os.path.join(workspace, 'INSTALLATION_AND_RUN_GUIDE.md')
    if os.path.exists(doc2_md_path):
        with open(doc2_md_path, 'r', encoding='utf-8') as f:
            md2 = f.read()
        elements2 = parse_markdown_to_structure(md2)
        create_word_doc(elements2, "AI-NIDS Installation & Run Guide", os.path.join(workspace, 'INSTALLATION_AND_RUN_GUIDE.docx'))
        create_pdf_doc(elements2, "AI-NIDS Installation & Run Guide", os.path.join(workspace, 'INSTALLATION_AND_RUN_GUIDE.pdf'))

if __name__ == '__main__':
    main()
